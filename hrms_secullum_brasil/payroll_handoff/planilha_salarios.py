"""Planilha de salários (docs/PRD-importacao-secullum.md, IS-13; Etapa 14).

CSV ou Excel com CPF, salário-base (ou valor da hora, para horista) e início → Salary Structure Assignment de cada
colaborador na Salary Structure "Padrão Ponto - <sigla>" da empresa (só o salário-base; as rubricas do ponto vão
por Additional Salary no Fechamento de Ponto). Linha com erro não para as outras; repetir a mesma planilha não
duplica.
"""

import csv
import io

import frappe
from frappe import _
from frappe.utils import flt, getdate
from frappe.utils.file_manager import get_file_path

from hrms_secullum_brasil.engine.documentos import somente_digitos
from hrms_secullum_brasil.engine.planilha_salarios import LinhaSalario, ler_planilha

PAPEIS = ("System Manager", "RH Ponto")
COMPONENTE = "Salário Base"
ABREVIACAO = "SB"
SSA = "Salary Structure Assignment"


def nome_estrutura(company: str) -> str:
	return f"Padrão Ponto - {frappe.db.get_value('Company', company, 'abbr')}"


def ler_arquivo(file_url: str) -> list[list]:
	"""Linhas da planilha (CSV com ; ou , — ou Excel .xlsx)."""
	caminho = get_file_path(file_url)
	if caminho.lower().endswith((".xlsx", ".xlsm")):
		from openpyxl import load_workbook

		planilha = load_workbook(caminho, read_only=True, data_only=True).active
		return [list(linha) for linha in planilha.iter_rows(values_only=True)]
	with open(caminho, "rb") as f:
		bruto = f.read()
	texto = (
		bruto.decode("utf-8-sig", errors="replace") if bruto[:3] == b"\xef\xbb\xbf" else _decodificar(bruto)
	)
	dialeto = csv.Sniffer().sniff(texto[:4096], delimiters=";,\t") if texto.strip() else csv.excel
	return [linha for linha in csv.reader(io.StringIO(texto), dialeto)]


def _decodificar(bruto: bytes) -> str:
	try:
		return bruto.decode("utf-8")
	except UnicodeDecodeError:
		return bruto.decode("latin-1")  # Excel em português salva CSV em Windows-1252


def garantir_estrutura(company: str) -> str:
	"""Salary Component "Salário Base" e a Salary Structure "Padrão Ponto - <sigla>" (submetida) da empresa."""
	if not frappe.db.exists("Salary Component", COMPONENTE):
		frappe.get_doc(
			{
				"doctype": "Salary Component",
				"salary_component": COMPONENTE,
				"salary_component_abbr": ABREVIACAO,
				"type": "Earning",
			}
		).insert(ignore_permissions=True)
	_garantir_conta_da_folha(company)
	nome = nome_estrutura(company)
	if not frappe.db.exists("Salary Structure", nome):
		estrutura = frappe.get_doc(
			{
				"doctype": "Salary Structure",
				"name": nome,
				"company": company,
				"payroll_frequency": "Monthly",
				"currency": frappe.db.get_value("Company", company, "default_currency") or "BRL",
				"earnings": [
					{
						"salary_component": COMPONENTE,
						"abbr": ABREVIACAO,
						"amount_based_on_formula": 1,
						"formula": "base",
					}
				],
			}
		)
		estrutura.insert(ignore_permissions=True)
		estrutura.submit()
	return nome


def _garantir_conta_da_folha(company: str) -> None:
	"""O HRMS exige a conta "Payroll Payable" da empresa (o plano de contas padrão já a cria) e, desde o v16, que ela
	tenha o tipo "Payable" (o Payroll Entry recusa o envio sem isso)."""
	conta = frappe.db.get_value("Company", company, "default_payroll_payable_account") or frappe.db.get_value(
		"Account", {"company": company, "account_name": "Payroll Payable", "is_group": 0}, "name"
	)
	if not conta:
		return
	frappe.db.set_value("Company", company, "default_payroll_payable_account", conta)
	if not frappe.db.get_value("Account", conta, "account_type"):
		frappe.db.set_value("Account", conta, "account_type", "Payable")


def _vinculo(linha: LinhaSalario) -> tuple[str | None, str | None]:
	"""(Employee, erro) do CPF: o vínculo ativo; com mais de um, o da empresa (CNPJ) da planilha."""
	ativos = frappe.get_all(
		"Employee",
		filters={"cpf": linha.cpf, "status": "Active"},
		fields=["name", "company", "date_of_joining"],
	)
	if linha.empresa:
		por_cnpj = {
			somente_digitos(c.tax_id): c.name for c in frappe.get_all("Company", fields=["name", "tax_id"])
		}
		ativos = [a for a in ativos if a.company == por_cnpj.get(linha.empresa)]
	if not ativos:
		return None, _("nenhum colaborador ativo com este CPF{0}").format(
			_(" nesta empresa") if linha.empresa else ""
		)
	if len(ativos) > 1:
		return None, _("o CPF tem vínculo ativo em mais de uma empresa: informe a coluna Empresa (CNPJ)")
	return ativos[0].name, None


@frappe.whitelist()
def importar(file_url: str, simular=0) -> dict:
	"""Importa a planilha. Devolve {criadas, sem_mudanca, erros: ["Linha N: ..."], avisos}."""
	frappe.only_for(PAPEIS)
	return importar_linhas(ler_arquivo(file_url), simular=bool(int(simular)))


def importar_linhas(linhas: list[list], simular: bool = False) -> dict:
	validas, erros = ler_planilha(linhas)
	resultado = {"criadas": 0, "sem_mudanca": 0, "substituidas": 0, "erros": erros, "avisos": []}
	for linha in validas:
		employee, erro = _vinculo(linha)
		if erro:
			resultado["erros"].append(f"Linha {linha.linha}: {erro}.")
			continue
		frappe.db.savepoint("planilha_salario")
		try:
			situacao = _gravar(employee, linha, resultado["avisos"], simular)
		except Exception as exc:
			frappe.db.rollback(save_point="planilha_salario")
			resultado["erros"].append(f"Linha {linha.linha}: {exc}")
			continue
		resultado[situacao] += 1
	return resultado


def _gravar(employee: str, linha: LinhaSalario, avisos: list, simular: bool) -> str:
	emp = frappe.db.get_value(
		"Employee", employee, ["company", "date_of_joining", "employee_name"], as_dict=True
	)
	inicio = linha.inicio
	if emp.date_of_joining and inicio < getdate(emp.date_of_joining):
		inicio = getdate(emp.date_of_joining)  # o HRMS não aceita início antes da admissão
		avisos.append(
			f"Linha {linha.linha}: início antes da admissão de {emp.employee_name}; usado {inicio:%d/%m/%Y}."
		)
	existente = frappe.get_all(
		SSA,
		filters={"employee": employee, "docstatus": 1, "from_date": inicio},
		fields=["name", "base"],
		limit=1,
	)
	if existente and flt(existente[0].base) == flt(linha.valor):
		return "sem_mudanca"
	if simular:
		return "substituidas" if existente else "criadas"
	estrutura = garantir_estrutura(emp.company)
	if existente:
		# mesmo início com outro valor: a planilha corrige o salário (cancela a atribuição anterior)
		frappe.get_doc(SSA, existente[0].name).cancel()
	ssa = frappe.get_doc(
		{
			"doctype": SSA,
			"employee": employee,
			"salary_structure": estrutura,
			"company": emp.company,
			"from_date": inicio,
			"base": linha.valor,
			"currency": frappe.db.get_value("Company", emp.company, "default_currency") or "BRL",
		}
	)
	ssa.insert(ignore_permissions=True)
	ssa.submit()
	return "substituidas" if existente else "criadas"


@frappe.whitelist()
def modelo() -> str:
	"""CSV modelo com os ativos importados sem salário (CPF, nome, cargo, tipo, salário vazio, início)."""
	frappe.only_for(PAPEIS)
	com_salario = set(frappe.get_all(SSA, filters={"docstatus": 1}, pluck="employee", distinct=True))
	saida = io.StringIO()
	escritor = csv.writer(saida, delimiter=";")
	escritor.writerow(
		["CPF", "Nome", "Cargo", "Tipo de remuneração", "Salário (horista: valor da hora)", "Início", "CNPJ"]
	)
	cnpjs = dict(frappe.get_all("Company", fields=["name", "tax_id"], as_list=True))
	for e in frappe.get_all(
		"Employee",
		filters={"status": "Active", "cpf": ["is", "set"]},
		fields=[
			"name",
			"cpf",
			"employee_name",
			"designation",
			"tipo_remuneracao",
			"company",
			"date_of_joining",
		],
		order_by="employee_name",
	):
		if e.name in com_salario:
			continue
		escritor.writerow(
			[
				e.cpf,
				e.employee_name,
				e.designation or "",
				e.tipo_remuneracao or "Mensalista",
				"",
				getdate(e.date_of_joining).strftime("%d/%m/%Y") if e.date_of_joining else "",
				somente_digitos(cnpjs.get(e.company)),
			]
		)
	return saida.getvalue()
